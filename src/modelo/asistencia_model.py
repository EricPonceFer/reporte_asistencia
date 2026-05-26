import pandas as pd
from pathlib import Path
from config.config import (
    RUTA_CODIGOS,
    RUTA_VELADA
)
from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Font,
    Border,
    Side
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from datetime import datetime, timedelta

class ErrorArchivo(Exception):
    """Errores relacionados con archivos."""
    pass


class ErrorProcesamiento(Exception):
    """Errores relacionados con el procesamiento."""
    pass


class ValidadorArchivos:

    COLUMNAS_ASISTENCIA = [
        "Codigo",
        "Fecha y hora"
    ]

    COLUMNAS_CODIGOS = [
        "CODIGO",
        "NOMBRES"
    ]

    COLUMNAS_VELADA = [
        "CODIGO",
        "Fecha Inicio",
        "Hora de ingreso",
        "Fecha Finalización"
    ]

    @staticmethod
    def validar_existencia_archivo(ruta):

        if not Path(ruta).exists():
            raise ErrorArchivo(
                f"El archivo no existe:\n{ruta}"
            )

    @staticmethod
    def validar_extension_excel(ruta):

        extensiones_validas = [".xlsx", ".xls"]

        if Path(ruta).suffix.lower() not in extensiones_validas:
            raise ErrorArchivo(
                "El archivo seleccionado no es un Excel válido."
            )

    @staticmethod
    def validar_columnas(dataframe, columnas_requeridas, nombre_archivo):

        columnas_faltantes = [
            columna
            for columna in columnas_requeridas
            if columna not in dataframe.columns
        ]

        if columnas_faltantes:
            raise ErrorArchivo(
                f"El archivo '{nombre_archivo}' no contiene las columnas necesarias:\n"
                f"{', '.join(columnas_faltantes)}"
            )


class AsistenciaModel:

    def __init__(
            self,
            ruta_asistencia,
            ruta_guardado
        ):

            # ==========================================
            # RUTAS USUARIO
            # ==========================================

            self.ruta_asistencia = Path(
                ruta_asistencia
            )

            self.ruta_guardado = Path(
                ruta_guardado
            )

            # ==========================================
            # RUTAS INTERNAS
            # ==========================================

            self.ruta_codigos = RUTA_CODIGOS

            self.ruta_velada = RUTA_VELADA

            # ==========================================
            # DATAFRAMES
            # ==========================================

            self.dataframe_asistencia = None

            self.dataframe_codigos = None

            self.dataframe_velada = None

        # =========================================================
        # CARGA DE ARCHIVOS
        # =========================================================

    def cargar_archivos(self):

        try:

            self._validar_archivos()

            # ==========================================
            # ASISTENCIA
            # ==========================================

            self.dataframe_asistencia = pd.read_excel(
                self.ruta_asistencia
            )

            # ==========================================
            # CODIGOS
            # ==========================================

            self.dataframe_codigos = pd.read_csv(
                self.ruta_codigos
            )

            # ==========================================
            # VELADA
            # ==========================================

            self.dataframe_velada = pd.read_csv(
                self.ruta_velada
            )
            self._limpiar_columnas()
            return True
        except Exception as error:
            raise ErrorArchivo(str(error))

    def _validar_archivos(self):

        # ==========================================
        # ASISTENCIA
        # ==========================================

        ValidadorArchivos.validar_existencia_archivo(
            self.ruta_asistencia
        )

        ValidadorArchivos.validar_extension_excel(
            self.ruta_asistencia
        )

        # ==========================================
        # ARCHIVOS INTERNOS
        # ==========================================

        ValidadorArchivos.validar_existencia_archivo(
            self.ruta_codigos
        )

        ValidadorArchivos.validar_existencia_archivo(
            self.ruta_velada
        )

    def _limpiar_columnas(self):

        self.dataframe_asistencia.columns = (
            self.dataframe_asistencia.columns.str.strip()
        )

        self.dataframe_codigos.columns = (
            self.dataframe_codigos.columns.str.strip()
        )

        self.dataframe_velada.columns = (
            self.dataframe_velada.columns.str.strip()
        )

        ValidadorArchivos.validar_columnas(
            self.dataframe_asistencia,
            ValidadorArchivos.COLUMNAS_ASISTENCIA,
            "Asistencia"
        )

        ValidadorArchivos.validar_columnas(
            self.dataframe_codigos,
            ValidadorArchivos.COLUMNAS_CODIGOS,
            "Código Barras"
        )

        ValidadorArchivos.validar_columnas(
            self.dataframe_velada,
            ValidadorArchivos.COLUMNAS_VELADA,
            "Personal de velada"
        )

    def _filtrar_rango_fechas(
        self,
        fecha_inicio,
        fecha_final
    ):
        try:

            fecha_inicio = pd.to_datetime(
                fecha_inicio
            )

            fecha_final = pd.to_datetime(
                fecha_final
            )

            if fecha_inicio > fecha_final:

                raise ErrorProcesamiento(
                    "La fecha inicial no puede ser mayor a la final."
                )

            mascara = (
                self.dataframe_asistencia["Fecha y hora"]
                .dt.date
                .between(
                    fecha_inicio.date(),
                    fecha_final.date()
                )
            )

            self.dataframe_asistencia = (
                self.dataframe_asistencia[
                    mascara
                ].copy()
            )

            if self.dataframe_asistencia.empty:

                raise ErrorProcesamiento(
                    "No existen registros dentro del rango seleccionado."
                )

        except Exception as error:

            raise ErrorProcesamiento(
                f"Error filtrando fechas:\n{error}"
            )
    # =========================================================
    # PROCESAMIENTO PRINCIPAL
    # =========================================================

    def procesar_asistencia(self, fecha_inicio, fecha_final):

        try:

            self._convertir_fechas()
            self._filtrar_rango_fechas(
                fecha_inicio,
                fecha_final
            )
            self._crear_columnas_fecha()
            self._unir_codigos_nombres()
            self._procesar_turnos_velada()
            self._identificar_personal_velada()

            dataframe_normal = self._procesar_personal_normal()
            dataframe_velada = self._procesar_personal_velada()

            dataframe_final = pd.concat(
                [dataframe_normal, dataframe_velada],
                ignore_index=True
            )

            dataframe_final = dataframe_final.sort_values(
                by=["Fecha", "Nombre Completo"]
            )

            return dataframe_final

        except Exception as error:
            raise ErrorProcesamiento(
                f"Ocurrió un error durante el procesamiento:\n{error}"
            )

    # =========================================================
    # PREPARACIÓN DE DATOS
    # =========================================================

    def _convertir_fechas(self):

        self.dataframe_asistencia["Fecha y hora"] = pd.to_datetime(
            self.dataframe_asistencia["Fecha y hora"],
            errors="coerce"
        )

        if self.dataframe_asistencia["Fecha y hora"].isna().all():
            raise ErrorProcesamiento(
                "No fue posible convertir las fechas del archivo de asistencia."
            )

    def _crear_columnas_fecha(self):

        fecha = self.dataframe_asistencia["Fecha y hora"]

        self.dataframe_asistencia["dia"] = fecha.dt.day
        self.dataframe_asistencia["mes"] = fecha.dt.month
        self.dataframe_asistencia["año"] = fecha.dt.year
        self.dataframe_asistencia["semana"] = (
            fecha.dt.isocalendar().week
        )

    def _unir_codigos_nombres(self):

        self.dataframe_asistencia = pd.merge(
            self.dataframe_asistencia,
            self.dataframe_codigos[["CODIGO", "NOMBRES"]],
            left_on="Codigo",
            right_on="CODIGO",
            how="left"
        )

        self.dataframe_asistencia.drop(
            columns=["CODIGO"],
            inplace=True
        )

    def _procesar_turnos_velada(self):

        self.dataframe_velada["fecha inicio y hora"] = pd.to_datetime(
            self.dataframe_velada["Fecha Inicio"].astype(str)
            + " "
            + self.dataframe_velada["Hora de ingreso"].astype(str),
            errors="coerce"
        )

        self.dataframe_velada["Fecha Finalización"] = pd.to_datetime(
            self.dataframe_velada["Fecha Finalización"].astype(str)
            + " 07:00:00",
            errors="coerce"
        )

    def _identificar_personal_velada(self):

        self.dataframe_asistencia["velada"] = (
            self.dataframe_asistencia.apply(
                lambda fila: self._esta_en_turno(fila),
                axis=1
            )
        )

    def _esta_en_turno(self, fila):

        turnos_operador = self.dataframe_velada[
            self.dataframe_velada["CODIGO"] == fila["Codigo"]
        ]

        return any(
            inicio <= fila["Fecha y hora"] <= fin
            for inicio, fin in zip(
                turnos_operador["fecha inicio y hora"],
                turnos_operador["Fecha Finalización"]
            )
        )

    # =========================================================
    # PERSONAL NORMAL
    # =========================================================

    def _procesar_personal_normal(self):

        dataframe_normal = self.dataframe_asistencia[
            ~self.dataframe_asistencia["velada"]
        ].copy()

        resultados = []

        for codigo_operador, contenido_operador in dataframe_normal.groupby("Codigo"):

            contenido_operador = contenido_operador.sort_values(
                by="Fecha y hora"
            )

            contenido_operador["fecha"] = (
                contenido_operador["Fecha y hora"].dt.date
            )

            for _, data_dia in contenido_operador.groupby("fecha"):

                resultado = self._crear_registro_personal_normal(
                    codigo_operador,
                    data_dia
                )

                resultados.append(resultado)

        return pd.DataFrame(resultados)

    def _crear_registro_personal_normal(self, codigo_operador, data_dia):

        hora_entrada = data_dia.iloc[0]["Fecha y hora"]
        hora_salida = data_dia.iloc[-1]["Fecha y hora"]

        horas_trabajadas = (
            hora_salida - hora_entrada
        ).total_seconds() / 3600

        if horas_trabajadas >= 14 or horas_trabajadas < 1:
            hora_salida = pd.NaT
            horas_trabajadas = 0

        return {
            "Nombre Completo": data_dia.iloc[0]["NOMBRES"],
            "Código": codigo_operador,
            "Fecha": hora_entrada.date(),
            "Año": hora_entrada.year,
            "Mes": hora_entrada.month,
            "Semana": hora_entrada.isocalendar().week,
            "Día": hora_entrada.day,
            "Hora_entrada": hora_entrada.time(),
            "Hora_salida": (
                hora_salida.time()
                if pd.notna(hora_salida)
                else None
            ),
            "Horas_trabajadas": round(horas_trabajadas, 2)
        }

    # =========================================================
    # PERSONAL VELADA
    # =========================================================

    def _procesar_personal_velada(self):

        dataframe_velada = self.dataframe_asistencia[
            self.dataframe_asistencia["velada"]
        ].copy()

        resultados = []

        for codigo_operador, grupo in dataframe_velada.groupby("Codigo"):

            grupo = grupo.sort_values(
                "Fecha y hora"
            ).reset_index(drop=True)

            tiempos = grupo["Fecha y hora"].tolist()

            i = 0
            n = len(tiempos)

            while i < n:

                inicio = tiempos[i]

                if inicio.hour >= 17:
                    fecha_base = inicio.date() + pd.Timedelta(days=1)
                else:

                    resultados.append(
                        self._registro_velada_sin_entrada(
                            grupo,
                            codigo_operador,
                            inicio
                        )
                    )

                    i += 1
                    continue

                if i + 1 < n:

                    fin = tiempos[i + 1]

                    duracion = (
                        fin - inicio
                    ).total_seconds() / 3600

                    if 5 < duracion <= 14:

                        resultados.append(
                            self._registro_velada_completo(
                                grupo,
                                codigo_operador,
                                inicio,
                                fin,
                                fecha_base,
                                duracion
                            )
                        )

                        i += 2

                    else:

                        resultados.append(
                            self._registro_velada_sin_salida(
                                grupo,
                                codigo_operador,
                                inicio,
                                fecha_base
                            )
                        )

                        i += 1

                else:

                    resultados.append(
                        self._registro_velada_sin_salida(
                            grupo,
                            codigo_operador,
                            inicio,
                            fecha_base
                        )
                    )

                    i += 1

        return pd.DataFrame(resultados)

    # =========================================================
    # REGISTROS VELADA
    # =========================================================

    def _registro_velada_completo(
        self,
        grupo,
        codigo_operador,
        inicio,
        fin,
        fecha_base,
        duracion
    ):

        return {
            "Nombre Completo": grupo.iloc[0]["NOMBRES"],
            "Código": codigo_operador,
            "Fecha": fecha_base,
            "Año": fecha_base.year,
            "Mes": fecha_base.month,
            "Semana": fecha_base.isocalendar().week,
            "Día": fecha_base.day,
            "Hora_entrada": inicio.time(),
            "Hora_salida": fin.time(),
            "Horas_trabajadas": round(duracion, 2)
        }

    def _registro_velada_sin_salida(
        self,
        grupo,
        codigo_operador,
        inicio,
        fecha_base
    ):

        return {
            "Nombre Completo": grupo.iloc[0]["NOMBRES"],
            "Código": codigo_operador,
            "Fecha": fecha_base,
            "Año": fecha_base.year,
            "Mes": fecha_base.month,
            "Semana": fecha_base.isocalendar().week,
            "Día": fecha_base.day,
            "Hora_entrada": inicio.time(),
            "Hora_salida": None,
            "Horas_trabajadas": 0
        }

    def _registro_velada_sin_entrada(
        self,
        grupo,
        codigo_operador,
        inicio
    ):

        fecha = inicio.date() + pd.Timedelta(days=1)

        return {
            "Nombre Completo": grupo.iloc[0]["NOMBRES"],
            "Código": codigo_operador,
            "Fecha": fecha,
            "Año": inicio.year,
            "Mes": inicio.month,
            "Semana": fecha.isocalendar().week,
            "Día": fecha.day,
            "Hora_entrada": None,
            "Hora_salida": inicio.time(),
            "Horas_trabajadas": 0
        }
    
    # =========================================================
    # Reporte En Excel
    # =========================================================
    def _generar_formato_semanal(
        self,
        dataframe,
        ruta_final
    ):

        try:

            if dataframe.empty:

                raise ErrorProcesamiento(
                    "El DataFrame está vacío."
                )

            # =====================================================
            # ESTILOS
            # =====================================================

            borde_fino = Border(
                left=Side(style="thin"),
                right=Side(style="thin"),
                top=Side(style="thin"),
                bottom=Side(style="thin")
            )

            # =====================================================
            # WORKBOOK
            # =====================================================

            wb = Workbook()

            wb.remove(wb.active)

            # =====================================================
            # DÍAS Y MESES
            # =====================================================

            dias_semana = [
                "LUNES",
                "MARTES",
                "MIÉRCOLES",
                "JUEVES",
                "VIERNES",
                "SÁBADO",
                "DOMINGO"
            ]

            meses_es = {
                1: "ENERO",
                2: "FEBRERO",
                3: "MARZO",
                4: "ABRIL",
                5: "MAYO",
                6: "JUNIO",
                7: "JULIO",
                8: "AGOSTO",
                9: "SEPTIEMBRE",
                10: "OCTUBRE",
                11: "NOVIEMBRE",
                12: "DICIEMBRE"
            }

            # =====================================================
            # OPERADORES ACTIVOS
            # =====================================================

            operadores_activos = (
                self.dataframe_codigos[
                    self.dataframe_codigos["ACTIVO"] == "SI"
                ]["NOMBRES"]
                .dropna()
                .unique()
            )

            operadores_activos = sorted(
                operadores_activos
            )

            # =====================================================
            # AGRUPAR POR AÑO Y SEMANA
            # =====================================================

            for (anio, semana), grupo_semana in dataframe.groupby(
                ["Año", "Semana"]
            ):

                # =================================================
                # HOJA
                # =================================================

                ws = wb.create_sheet(
                    title=f"{anio}-S{semana}"
                )

                # =================================================
                # CONFIGURACIÓN IMPRESIÓN
                # =================================================

                ws.page_setup.orientation = (
                    ws.ORIENTATION_LANDSCAPE
                )

                ws.page_setup.fitToWidth = 1

                ws.page_setup.fitToHeight = False

                ws.page_margins = PageMargins(
                    left=0.2,
                    right=0.2,
                    top=0.3,
                    bottom=0.3
                )

                ws.print_options.horizontalCentered = True

                # =================================================
                # FECHAS
                # =================================================

                mes = int(
                    grupo_semana["Mes"].iloc[0]
                )

                nombre_mes = meses_es.get(
                    mes,
                    ""
                )

                lunes_semana = datetime.fromisocalendar(
                    int(anio),
                    int(semana),
                    1
                )

                # =================================================
                # TÍTULO
                # =================================================

                ws.merge_cells(
                    start_row=1,
                    start_column=2,
                    end_row=1,
                    end_column=15
                )

                ws["B1"] = (
                    f"SEMANA {semana} - "
                    f"{nombre_mes} {anio}"
                )

                ws["B1"].alignment = Alignment(
                    horizontal="center"
                )

                ws["B1"].font = Font(
                    bold=True,
                    size=14
                )

                # Bordes título
                for col in range(2, 16):

                    ws.cell(
                        row=1,
                        column=col
                    ).border = borde_fino

                # =================================================
                # ENCABEZADO IZQUIERDO
                # =================================================

                ws["A2"] = f"MES: {nombre_mes}"

                ws["A2"].font = Font(
                    bold=True
                )

                ws.merge_cells("A3:A4")

                ws["A3"] = "OPERADOR"

                ws["A3"].alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

                ws["A3"].font = Font(
                    bold=True
                )

                for fila in range(2, 5):

                    ws.cell(
                        row=fila,
                        column=1
                    ).border = borde_fino

                # =================================================
                # ENCABEZADO DÍAS
                # =================================================

                col = 2

                for i, dia_nombre in enumerate(
                    dias_semana
                ):

                    fecha_real = (
                        lunes_semana
                        + timedelta(days=i)
                    )

                    numero_dia = fecha_real.day

                    # =============================================
                    # DÍA
                    # =============================================

                    ws.merge_cells(
                        start_row=2,
                        start_column=col,
                        end_row=2,
                        end_column=col + 1
                    )

                    ws.cell(
                        row=2,
                        column=col
                    ).value = dia_nombre

                    ws.cell(
                        row=2,
                        column=col
                    ).alignment = Alignment(
                        horizontal="center"
                    )

                    ws.cell(
                        row=2,
                        column=col
                    ).font = Font(
                        bold=True
                    )

                    # =============================================
                    # NÚMERO DÍA
                    # =============================================

                    ws.merge_cells(
                        start_row=3,
                        start_column=col,
                        end_row=3,
                        end_column=col + 1
                    )

                    ws.cell(
                        row=3,
                        column=col
                    ).value = numero_dia

                    ws.cell(
                        row=3,
                        column=col
                    ).alignment = Alignment(
                        horizontal="center"
                    )

                    # =============================================
                    # ENTRADA / SALIDA
                    # =============================================

                    ws.cell(
                        row=4,
                        column=col
                    ).value = "ENTRADA"

                    ws.cell(
                        row=4,
                        column=col + 1
                    ).value = "SALIDA"

                    ws.cell(
                        row=4,
                        column=col
                    ).alignment = Alignment(
                        horizontal="center"
                    )

                    ws.cell(
                        row=4,
                        column=col + 1
                    ).alignment = Alignment(
                        horizontal="center"
                    )

                    # =============================================
                    # BORDES
                    # =============================================

                    for fila in range(2, 5):

                        for columna in range(
                            col,
                            col + 2
                        ):

                            ws.cell(
                                row=fila,
                                column=columna
                            ).border = borde_fino

                    col += 2

                # =================================================
                # DATOS
                # =================================================

                fila_excel = 5

                for nombre in operadores_activos:

                    grupo_operador = grupo_semana[
                        grupo_semana["Nombre Completo"]
                        == nombre
                    ]

                    # =============================================
                    # NOMBRE OPERADOR
                    # =============================================

                    ws.cell(
                        row=fila_excel,
                        column=1
                    ).value = nombre

                    ws.cell(
                        row=fila_excel,
                        column=1
                    ).alignment = Alignment(
                        horizontal="left"
                    )

                    ws.cell(
                        row=fila_excel,
                        column=1
                    ).border = borde_fino

                    # =============================================
                    # DÍAS
                    # =============================================

                    col = 2

                    for i in range(7):

                        fecha_real = (
                            lunes_semana
                            + timedelta(days=i)
                        ).date()

                        registro = grupo_operador[
                            grupo_operador["Fecha"]
                            == fecha_real
                        ]

                        if not registro.empty:

                            hora_entrada = registro.iloc[0][
                                "Hora_entrada"
                            ]

                            hora_salida = registro.iloc[0][
                                "Hora_salida"
                            ]

                            # =====================================
                            # HORA ENTRADA
                            # =====================================

                            if pd.notna(hora_entrada):

                                celda_entrada = ws.cell(
                                    row=fila_excel,
                                    column=col
                                )

                                celda_entrada.value = (
                                    hora_entrada
                                )

                                celda_entrada.number_format = (
                                    "hh:mm"
                                )

                            # =====================================
                            # HORA SALIDA
                            # =====================================

                            if pd.notna(hora_salida):

                                celda_salida = ws.cell(
                                    row=fila_excel,
                                    column=col + 1
                                )

                                celda_salida.value = (
                                    hora_salida
                                )

                                celda_salida.number_format = (
                                    "hh:mm"
                                )

                        # =========================================
                        # ESTILOS CELDAS
                        # =========================================

                        ws.cell(
                            row=fila_excel,
                            column=col
                        ).border = borde_fino

                        ws.cell(
                            row=fila_excel,
                            column=col + 1
                        ).border = borde_fino

                        ws.cell(
                            row=fila_excel,
                            column=col
                        ).alignment = Alignment(
                            horizontal="center"
                        )

                        ws.cell(
                            row=fila_excel,
                            column=col + 1
                        ).alignment = Alignment(
                            horizontal="center"
                        )

                        col += 2

                    fila_excel += 1

                # =================================================
                # ÁREA IMPRESIÓN
                # =================================================

                ws.print_area = (
                    f"A1:O{fila_excel - 1}"
                )

                # =================================================
                # AUTO WIDTH COLUMNA A
                # =================================================

                max_length = 0

                for fila in ws.iter_rows(
                    min_col=1,
                    max_col=1
                ):

                    for cell in fila:

                        if cell.value:

                            max_length = max(
                                max_length,
                                len(str(cell.value))
                            )

                ws.column_dimensions["A"].width = (
                    max_length + 6
                )

                # =================================================
                # ANCHO COLUMNAS
                # =================================================

                for i in range(2, 17):

                    ws.column_dimensions[
                        get_column_letter(i)
                    ].width = 11

            # =====================================================
            # GUARDAR
            # =====================================================

            wb.save(ruta_final)

        except PermissionError:

            raise ErrorProcesamiento(
                "Cierre el archivo Excel antes de guardar."
            )

        except Exception as error:

            raise ErrorProcesamiento(
                f"Error generando formato semanal:\n{error}"
            )

    # =========================================================
    # EXPORTAR
    # =========================================================

    def exportar_excel(
        self,
        dataframe,
        nombre_archivo
    ):

        try:

            ruta_final = (
                self.ruta_guardado
                / nombre_archivo
            )

            self._generar_formato_semanal(
                dataframe,
                ruta_final
            )

        except Exception as error:

            raise ErrorProcesamiento(
                f"No fue posible exportar el Excel:\n{error}"
            )